"""
Unit tests untuk fitur UAT Result (dokumen naratif per-case + Daftar Isi).
"""

import sys
import os
import io

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document

from main import (
    build_uat_result_document,
    _clean_case_title,
    _compress_headers_inline,
    _build_uat_result_case_text,
)


def _para_texts(doc):
    return [p.text for p in doc.paragraphs]


REMARKS_VA = (
    'URL :\n"https://x/balance"\n\n'
    "Header :[\nContent-Type=application/json\n\"Authorization=Bearer xyz\"\n]\n\n"
    'Request:\n{\n  "accountNo": "123"\n}\n\n'
    'Response :\n{\n  "responseCode": "2001100"\n}'
)


class TestCleanCaseTitle:
    def test_qris_buang_prefix_nomor(self):
        assert _clean_case_title("18,1 Access Token Invalid", "QRIS") == "Access Token Invalid"
        assert _clean_case_title("3,5 Cannot use X", "QRIS") == "Cannot use X"

    def test_ftva_biarkan(self):
        assert _clean_case_title("Access Token Invalid", "FT_VA") == "Access Token Invalid"


class TestCompress:
    def test_header_jadi_satu_baris(self):
        raw = "Content-Type=application/json\n\"Authorization=Bearer xyz\""
        out = _compress_headers_inline(raw)
        assert out == '[Content-Type=application/json, "Authorization=Bearer xyz"]'

    def test_body_dan_response_dicompress(self):
        lines = _build_uat_result_case_text(REMARKS_VA)
        joined = "\n".join(lines)
        # URL tetap ada barisnya
        assert "URL :" in joined
        # Header jadi satu baris
        assert 'Header :[Content-Type=application/json, "Authorization=Bearer xyz"]' in joined
        # Body & Response compress (tanpa newline dalam JSON)
        assert '{"accountNo":"123"}' in joined
        assert '{"responseCode":"2001100"}' in joined


class TestBuildUatResult:
    def _mini_uat(self):
        return {
            "Balance Services": [
                {"nomor_kasus_tes": "1.1", "langkah_tes": "Access Token Invalid",
                 "hasil_aktual": "Berhasil", "remarks": REMARKS_VA},
                {"nomor_kasus_tes": "1.7", "langkah_tes": "Inquiry lebih dari 1",
                 "hasil_aktual": "Tidak dites", "remarks": "Tidak dites karena X."},
            ],
        }

    def test_heading_section_dan_case(self):
        doc = build_uat_result_document(self._mini_uat(), title="UAT Result", product_key="FT_VA")
        h1 = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
        h2 = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
        assert "Balance Services" in h1
        assert "1.1 Access Token Invalid" in h2
        assert "1.7 Inquiry lebih dari 1" in h2

    def test_tidak_dites_tampil_alasan(self):
        doc = build_uat_result_document(self._mini_uat(), product_key="FT_VA")
        txt = "\n".join(_para_texts(doc))
        assert "Tidak dites karena X." in txt

    def test_ada_daftar_isi_toc(self):
        # Field TOC diselipkan -> cek ada instruksi TOC dalam XML dokumen
        doc = build_uat_result_document(self._mini_uat(), product_key="FT_VA")
        xml = doc.element.xml
        assert "TOC" in xml
